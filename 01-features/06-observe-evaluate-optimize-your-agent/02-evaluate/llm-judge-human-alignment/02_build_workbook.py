"""Phase 1: build the SME ground-truth workbook from the stored sessions.

The evaluation team fills the gray columns (case, session, intent, risk, selection reason).
Domain experts complete the green columns (expected outcome, required action, unacceptable
outcome, rationale, expected tool calls, and assertions) before any review starts.

By default the domain-expert columns are prefilled from data/sme_ground_truth.json so the
sample runs end to end. Pass --blank to produce the empty workbook you would send to SMEs.

Usage:
    python 02_build_workbook.py [--blank]

Output:
    output/ground_truth_workbook.xlsx
"""

import argparse

from common import SESSIONS_FILE, SME_GROUND_TRUTH_FILE, WORKBOOK_FILE, read_json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

# (header, source, width). "team" columns come from sessions.json, "sme" columns from domain experts.
COLUMNS = [
    ("case_id", "team", 10),
    ("session_id", "team", 30),
    ("intent", "team", 14),
    ("risk", "team", 9),
    ("selection_reason", "team", 40),
    ("user_request", "team", 45),
    ("expected_business_outcome", "sme", 50),
    ("required_action", "sme", 40),
    ("unacceptable_outcome", "sme", 40),
    ("rationale", "sme", 40),
    ("expected_tool_calls", "sme", 32),
    ("assertions", "sme", 60),
    ("approved_by", "sme", 14),
]

TEAM_FILL = PatternFill("solid", fgColor="E9EBED")
SME_FILL = PatternFill("solid", fgColor="D5F2EB")
HEADER_FONT = Font(bold=True, name="Helvetica")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--blank", action="store_true", help="Leave the domain-expert columns empty")
    args = parser.parse_args()

    sessions = read_json(SESSIONS_FILE)["sessions"]
    ground_truth = {} if args.blank else read_json(SME_GROUND_TRUTH_FILE)["cases"]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ground_truth"
    for index, (header, source, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=header)
        cell.font = HEADER_FONT
        cell.fill = TEAM_FILL if source == "team" else SME_FILL
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.freeze_panes = "B2"

    for row, session in enumerate(sessions, start=2):
        truth = ground_truth.get(session["case_id"], {})
        values = {
            "case_id": session["case_id"],
            "session_id": session["session_id"],
            "intent": session["intent"],
            "risk": session["risk"],
            "selection_reason": session["selection_reason"],
            "user_request": session["turns"][0]["content"],
            "expected_business_outcome": truth.get("expected_business_outcome"),
            "required_action": truth.get("required_action"),
            "unacceptable_outcome": truth.get("unacceptable_outcome"),
            "rationale": truth.get("rationale"),
            # Multi-value cells use one item per line so SMEs can edit them in Excel.
            "expected_tool_calls": "\n".join(truth.get("expected_tool_calls", [])) or None,
            "assertions": "\n".join(truth.get("assertions", [])) or None,
            "approved_by": "sme-lead" if truth else None,
        }
        for column, (header, source, _) in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row=row, column=column, value=values[header])
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if source == "sme":
                cell.fill = SME_FILL

    risk_column = next(i for i, c in enumerate(COLUMNS, start=1) if c[0] == "risk")
    letter = sheet.cell(row=1, column=risk_column).column_letter
    validation = DataValidation(type="list", formula1='"low,medium,high"', allow_blank=False)
    validation.add(f"{letter}2:{letter}{len(sessions) + 1}")
    sheet.add_data_validation(validation)

    WORKBOOK_FILE.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(WORKBOOK_FILE)
    state = "blank SME columns" if args.blank else "SME columns prefilled from data/sme_ground_truth.json"
    print(f"  Wrote {WORKBOOK_FILE.name} with {len(sessions)} cases ({state})")


if __name__ == "__main__":
    main()
