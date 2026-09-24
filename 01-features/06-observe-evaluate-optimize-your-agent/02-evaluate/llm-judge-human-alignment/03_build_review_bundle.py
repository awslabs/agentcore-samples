"""Phase 2: lock the SME ground truth and build the bundle the review app loads.

Reads the completed workbook, checks that every case has the required domain-expert
fields, and joins each row with its stored session. The ground truth is hashed so any
later edit is visible: reviews and judge runs record the version they were scored against.

Usage:
    python 03_build_review_bundle.py

Output:
    output/review_cases.json  - sessions plus locked ground truth, loaded by review_app.html
"""

import hashlib
import json
from datetime import datetime, timezone

from common import REVIEW_CASES_FILE, SESSIONS_FILE, WORKBOOK_FILE, read_json, write_json
from openpyxl import load_workbook

REQUIRED_FIELDS = ("expected_business_outcome", "required_action", "unacceptable_outcome", "assertions")
LIST_FIELDS = ("expected_tool_calls", "assertions")


def read_workbook() -> dict[str, dict]:
    if not WORKBOOK_FILE.exists():
        raise SystemExit(f"{WORKBOOK_FILE.name} not found. Run `python 02_build_workbook.py` first.")
    sheet = load_workbook(WORKBOOK_FILE, read_only=True)["ground_truth"]
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    cases = {}
    for values in rows:
        row: dict = dict(zip(headers, values))
        if not row.get("case_id"):
            continue
        for field in LIST_FIELDS:
            row[field] = [line.strip() for line in str(row.get(field) or "").splitlines() if line.strip()]
        cases[row["case_id"]] = row
    return cases


def main() -> None:
    sessions = read_json(SESSIONS_FILE)
    rows = read_workbook()

    missing = [
        f"{case_id}: {field}" for case_id, row in rows.items() for field in REQUIRED_FIELDS if not row.get(field)
    ]
    unknown = sorted(set(rows) - {session["case_id"] for session in sessions["sessions"]})
    if missing or unknown:
        details = "\n  ".join(missing + [f"{case_id}: no stored session" for case_id in unknown])
        raise SystemExit(f"Ground truth is not ready to lock:\n  {details}")

    cases = []
    for session in sessions["sessions"]:
        row = rows.get(session["case_id"])
        if row is None:
            raise SystemExit(f"{session['case_id']} is missing from the workbook.")
        if row["session_id"] != session["session_id"]:
            raise SystemExit(f"{session['case_id']}: workbook session_id does not match sessions.json.")
        ground_truth = {
            field: row.get(field)
            for field in (
                "expected_business_outcome",
                "required_action",
                "unacceptable_outcome",
                "rationale",
                "expected_tool_calls",
                "assertions",
            )
        }
        cases.append({**session, "ground_truth": ground_truth})

    canonical = json.dumps({case["case_id"]: case["ground_truth"] for case in cases}, sort_keys=True)
    version = hashlib.sha256(canonical.encode("utf-8")).hexdigest()[:12]
    write_json(
        REVIEW_CASES_FILE,
        {
            "schema_version": "1.0",
            "ground_truth_version": version,
            "locked_at": datetime.now(timezone.utc).isoformat(),
            "service_name": sessions["service_name"],
            "log_group": sessions["log_group"],
            "cases": cases,
        },
    )
    print(f"  Locked ground truth version {version} for {len(cases)} cases")


if __name__ == "__main__":
    main()
